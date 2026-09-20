"""Tkinter GUI for the Student Management System."""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import csv
import os
import sys
from datetime import datetime

try:
    from openpyxl import Workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from student import Student
from database import (
    get_all_students, add_student, update_student,
    delete_student, find_student_by_id, search_students_by_name,
    search_students_by_code, search_students_by_course,
    search_students_by_email, search_students_all_fields,
    get_course_statistics, authenticate_admin, backup_database,
    restore_database, initialize_database, has_admin_accounts,
    create_admin
)


class AdminSetupDialog:
    """Dialog for creating the first administrator account."""

    def __init__(self, parent):
        self.parent = parent
        self.result = False
        self.username = ""

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Student Management System - Administrator Setup")
        self.dialog.resizable(False, False)
        self.dialog.grab_set()
        self.dialog.geometry("430x430")

        self._build_form()
        self._center_dialog()

        self.dialog.protocol("WM_DELETE_WINDOW", self._on_exit)
        self.dialog.bind("<Escape>", lambda e: self._on_exit())
        self.dialog.wait_window()

    def _center_dialog(self):
        self.dialog.update_idletasks()
        w = self.dialog.winfo_width()
        h = self.dialog.winfo_height()
        sw = self.dialog.winfo_screenwidth()
        sh = self.dialog.winfo_screenheight()
        self.dialog.geometry(f"+{(sw - w) // 2}+{(sh - h) // 2}")

    def _build_form(self):
        main_frame = ttk.Frame(self.dialog, padding=30)
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)

        ttk.Label(
            main_frame,
            text="Administrator Setup",
            font=("Segoe UI", 18, "bold")
        ).grid(row=0, column=0, pady=(0, 5))

        ttk.Label(
            main_frame,
            text="Create the first administrator account",
            font=("Segoe UI", 10)
        ).grid(row=1, column=0, pady=(0, 20))

        ttk.Label(main_frame, text="Username:").grid(
            row=2, column=0, sticky=tk.W, pady=(0, 5)
        )
        self.username_var = tk.StringVar()
        self.username_entry = ttk.Entry(
            main_frame, textvariable=self.username_var, width=32
        )
        self.username_entry.grid(
            row=3, column=0, sticky=tk.EW, pady=(0, 15)
        )

        ttk.Label(main_frame, text="Password:").grid(
            row=4, column=0, sticky=tk.W, pady=(0, 5)
        )
        self.password_var = tk.StringVar()
        self.password_entry = ttk.Entry(
            main_frame, textvariable=self.password_var,
            width=32, show="•"
        )
        self.password_entry.grid(
            row=5, column=0, sticky=tk.EW, pady=(0, 15)
        )

        ttk.Label(main_frame, text="Confirm Password:").grid(
            row=6, column=0, sticky=tk.W, pady=(0, 5)
        )
        self.confirm_var = tk.StringVar()
        self.confirm_entry = ttk.Entry(
            main_frame, textvariable=self.confirm_var,
            width=32, show="•"
        )
        self.confirm_entry.grid(
            row=7, column=0, sticky=tk.EW, pady=(0, 10)
        )

        self.show_password = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            main_frame,
            text="Show passwords",
            variable=self.show_password,
            command=self._toggle_password
        ).grid(row=8, column=0, sticky=tk.W, pady=(0, 10))

        self.error_var = tk.StringVar(value="")
        ttk.Label(
            main_frame,
            textvariable=self.error_var,
            font=("Segoe UI", 9),
            foreground="red",
            wraplength=360
        ).grid(row=9, column=0, sticky=tk.EW, pady=(0, 10))

        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=10, column=0, sticky=tk.EW, pady=(10, 0))
        btn_frame.columnconfigure(0, weight=1)
        btn_frame.columnconfigure(1, weight=1)

        ttk.Button(
            btn_frame,
            text="Create Administrator",
            command=self._on_create,
            style="Action.TButton"
        ).grid(row=0, column=0, sticky=tk.EW, padx=(0, 5))

        ttk.Button(
            btn_frame,
            text="Exit",
            command=self._on_exit,
            style="Action.TButton"
        ).grid(row=0, column=1, sticky=tk.EW, padx=(5, 0))

        self.username_entry.focus()

    def _toggle_password(self):
        show = "" if self.show_password.get() else "•"
        self.password_entry.config(show=show)
        self.confirm_entry.config(show=show)

    def _on_create(self):
        username = self.username_var.get().strip()
        password = self.password_var.get()
        confirm = self.confirm_var.get()

        if not username:
            self.error_var.set("Please enter a username.")
            self.username_entry.focus()
            return

        if not password:
            self.error_var.set("Please enter a password.")
            self.password_entry.focus()
            return

        if password != confirm:
            self.error_var.set("Passwords do not match.")
            self.confirm_entry.focus()
            return

        from database import create_admin
        success, message = create_admin(username, password)

        if success:
            self.username = username
            self.result = True
            messagebox.showinfo("Administrator Setup", message, parent=self.dialog)
            self.dialog.destroy()
        else:
            self.error_var.set(message)

    def _on_exit(self):
        self.result = False
        self.dialog.destroy()


class LoginDialog:
    """Login dialog for admin authentication."""

    def __init__(self, parent):
        self.parent = parent
        self.result = False
        self.username = ""
        
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Student Management System - Login")
        self.dialog.resizable(False, False)
        # Don't use transient(parent) when parent may be withdrawn
        self.dialog.grab_set()
        self.dialog.geometry("400x350")
        
        self._build_form()
        self._center_dialog()
        
        self.dialog.protocol("WM_DELETE_WINDOW", self._on_exit)
        self.dialog.bind("<Return>", lambda e: self._on_login())
        self.dialog.bind("<Escape>", lambda e: self._on_exit())
        self.dialog.wait_window()

    def _center_dialog(self):
        self.dialog.update_idletasks()
        w = self.dialog.winfo_width()
        h = self.dialog.winfo_height()
        sw = self.dialog.winfo_screenwidth()
        sh = self.dialog.winfo_screenheight()
        self.dialog.geometry(f"+{(sw - w) // 2}+{(sh - h) // 2}")

    def _build_form(self):
        main_frame = ttk.Frame(self.dialog, padding=30)
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)

        # Title
        ttk.Label(main_frame, text="Student Management System", 
                 font=("Segoe UI", 18, "bold")).grid(row=0, column=0, pady=(0, 5))
        ttk.Label(main_frame, text="Secure Administrator Login", 
                 font=("Segoe UI", 10), foreground="gray").grid(row=1, column=0, pady=(0, 25))

        # Username
        ttk.Label(main_frame, text="Username:", font=("Segoe UI", 10)).grid(
            row=2, column=0, sticky=tk.W, pady=(0, 5))
        self.username_var = tk.StringVar()
        self.username_entry = ttk.Entry(main_frame, textvariable=self.username_var, 
                                        width=30, font=("Segoe UI", 10))
        self.username_entry.grid(row=3, column=0, sticky=tk.EW, pady=(0, 15))
        self.username_entry.focus()

        # Password
        ttk.Label(main_frame, text="Password:", font=("Segoe UI", 10)).grid(
            row=4, column=0, sticky=tk.W, pady=(0, 5))
        
        pwd_frame = ttk.Frame(main_frame)
        pwd_frame.grid(row=5, column=0, sticky=tk.EW, pady=(0, 10))
        pwd_frame.columnconfigure(0, weight=1)
        
        self.password_var = tk.StringVar()
        self.password_entry = ttk.Entry(pwd_frame, textvariable=self.password_var, 
                                        width=30, font=("Segoe UI", 10), show="•")
        self.password_entry.grid(row=0, column=0, sticky=tk.EW)
        
        self.show_password = tk.BooleanVar(value=False)
        self.show_pwd_check = ttk.Checkbutton(pwd_frame, text="Show", 
                                              variable=self.show_password,
                                              command=self._toggle_password)
        self.show_pwd_check.grid(row=0, column=1, padx=(10, 0))

        # Error label (initially hidden)
        self.error_var = tk.StringVar(value="")
        self.error_label = ttk.Label(main_frame, textvariable=self.error_var, 
                                     font=("Segoe UI", 9), foreground="red")
        self.error_label.grid(row=6, column=0, sticky=tk.EW, pady=(5, 10))

        # Buttons
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=7, column=0, sticky=tk.EW, pady=(15, 0))
        btn_frame.columnconfigure(0, weight=1)
        btn_frame.columnconfigure(1, weight=1)
        
        ttk.Button(btn_frame, text="Login", command=self._on_login, 
                  style="Action.TButton").grid(row=0, column=0, sticky=tk.EW, padx=(0, 5))
        ttk.Button(btn_frame, text="Exit", command=self._on_exit, 
                  style="Action.TButton").grid(row=0, column=1, sticky=tk.EW, padx=(5, 0))

        # Bind Enter key in password field
        self.password_entry.bind("<Return>", lambda e: self._on_login())

    def _toggle_password(self):
        if self.show_password.get():
            self.password_entry.config(show="")
        else:
            self.password_entry.config(show="•")

    def _on_login(self):
        username = self.username_var.get().strip()
        password = self.password_var.get()
        
        if not username:
            self.error_var.set("Please enter username.")
            self.username_entry.focus()
            return
        if not password:
            self.error_var.set("Please enter password.")
            self.password_entry.focus()
            return
        
        self.error_var.set("")
        
        if authenticate_admin(username, password):
            self.username = username
            self.result = True
            self.dialog.destroy()
        else:
            self.error_var.set("Invalid username or password.")
            self.password_var.set("")
            self.password_entry.focus()

    def _on_exit(self):
        self.result = False
        self.dialog.destroy()


class BackupRestoreDialog:
    """Dialog for database backup and restore operations."""

    def __init__(self, parent):
        self.parent = parent
        
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Backup / Restore Database")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()
        self.dialog.geometry("450x450")
        
        self._build_dialog()
        self._center_dialog()
        
        self.dialog.protocol("WM_DELETE_WINDOW", self._on_close)
        self.dialog.bind("<Escape>", lambda e: self._on_close())
        self.dialog.wait_window()

    def _center_dialog(self):
        self.dialog.update_idletasks()
        w = self.dialog.winfo_width()
        h = self.dialog.winfo_height()
        sw = self.dialog.winfo_screenwidth()
        sh = self.dialog.winfo_screenheight()
        self.dialog.geometry(f"+{(sw - w) // 2}+{(sh - h) // 2}")

    def _build_dialog(self):
        main_frame = ttk.Frame(self.dialog, padding=25)
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)

        # Title
        ttk.Label(main_frame, text="Backup / Restore Database", 
                 font=("Segoe UI", 14, "bold")).grid(row=0, column=0, pady=(0, 10))
        ttk.Label(main_frame, text="Manage database backups and restoration", 
                 font=("Segoe UI", 9), foreground="gray").grid(row=1, column=0, pady=(0, 20))

        # Backup section
        backup_frame = ttk.LabelFrame(main_frame, text="Backup Database", padding=15)
        backup_frame.grid(row=2, column=0, sticky=tk.EW, pady=(0, 10))
        backup_frame.columnconfigure(0, weight=1)
        
        ttk.Label(backup_frame, text="Create a complete backup of the database including all students and admin data.", 
                 font=("Segoe UI", 9), wraplength=350).grid(row=0, column=0, sticky=tk.W, pady=(0, 10))
        
        ttk.Button(backup_frame, text="Create Backup", command=self._on_backup, 
                  style="Action.TButton").grid(row=1, column=0, sticky=tk.W)

        # Restore section
        restore_frame = ttk.LabelFrame(main_frame, text="Restore Database", padding=15)
        restore_frame.grid(row=3, column=0, sticky=tk.EW, pady=(0, 10))
        restore_frame.columnconfigure(0, weight=1)
        
        ttk.Label(restore_frame, text="Restore database from a backup file. This will replace all current data.", 
                 font=("Segoe UI", 9), wraplength=350).grid(row=0, column=0, sticky=tk.W, pady=(0, 10))
        
        ttk.Button(restore_frame, text="Restore from Backup", command=self._on_restore, 
                  style="Action.TButton").grid(row=1, column=0, sticky=tk.W)

        # Close button
        btn_frame = ttk.Frame(main_frame)
        btn_frame.grid(row=4, column=0, sticky=tk.EW, pady=(15, 0))
        ttk.Button(btn_frame, text="Close", command=self._on_close).pack(side=tk.RIGHT)

    def _on_backup(self):
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        default_name = f"student_management_backup_{timestamp}.db"
        
        filename = filedialog.asksaveasfilename(
            title="Save Database Backup",
            defaultextension=".db",
            initialfile=default_name,
            filetypes=[("SQLite Database", "*.db"), ("All Files", "*.*")]
        )
        if not filename:
            return
        
        success, message = backup_database(filename)
        if success:
            messagebox.showinfo("Success", message)
        else:
            messagebox.showerror("Backup Failed", message)

    def _on_restore(self):
        filename = filedialog.askopenfilename(
            title="Select Backup File to Restore",
            filetypes=[("SQLite Database", "*.db"), ("All Files", "*.*")]
        )
        if not filename:
            return
        
        # Confirmation dialog
        confirmed = messagebox.askyesno(
            "Confirm Restore",
            "This will replace the current database with the selected backup.\n"
            "Your current data will be backed up automatically before restoring.\n\n"
            "Continue with restore?",
            icon="warning"
        )
        if not confirmed:
            return
        
        success, message = restore_database(filename)
        if success:
            messagebox.showinfo("Success", message)
            # Signal parent to refresh
            self.parent.event_generate("<<DatabaseRestored>>")
        else:
            messagebox.showerror("Restore Failed", message)

    def _on_close(self):
        self.dialog.destroy()


class StudentFormDialog:
    """Modal dialog for adding or updating a student."""

    def __init__(self, parent, mode="add", student_data=None):
        self.mode = mode
        self.student_data = student_data
        self.result = None
        self.parent = parent

        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Add Student" if mode == "add" else "Update Student")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        self._build_form()
        if mode == "update" and student_data:
            self._populate(student_data)

        self.dialog.protocol("WM_DELETE_WINDOW", self._on_cancel)
        self.dialog.bind("<Return>", lambda e: self._on_save())
        self.dialog.bind("<Escape>", lambda e: self._on_cancel())
        self.dialog.wait_window()

    def _build_form(self):
        frame = ttk.Frame(self.dialog, padding=20)
        frame.pack(fill=tk.BOTH)

        self.fields = {}
        labels = ["Student Code", "Name", "Age", "Course", "Email"]
        for i, label in enumerate(labels):
            ttk.Label(frame, text=label + ":", font=("Segoe UI", 10)).grid(
                row=i, column=0, sticky=tk.W, pady=4
            )
            entry = ttk.Entry(frame, width=35, font=("Segoe UI", 10))
            entry.grid(row=i, column=1, padx=10, pady=4)
            self.fields[label.lower().replace(" ", "_")] = entry

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=len(labels), column=0, columnspan=2, pady=15)
        ttk.Button(btn_frame, text="Save", command=self._on_save).pack(
            side=tk.LEFT, padx=5
        )
        ttk.Button(btn_frame, text="Cancel", command=self._on_cancel).pack(
            side=tk.LEFT, padx=5
        )

    def _populate(self, data):
        self.fields["student_code"].insert(0, data["student_code"])
        self.fields["student_code"].config(state="readonly")
        self.fields["name"].insert(0, data["name"])
        self.fields["age"].insert(0, str(data["age"]))
        self.fields["course"].insert(0, data["course"])
        self.fields["email"].insert(0, data["email"])

    def _on_save(self):
        student_code = self.fields["student_code"].get().strip()
        name = self.fields["name"].get().strip()
        age_str = self.fields["age"].get().strip()
        course = self.fields["course"].get().strip()
        email = self.fields["email"].get().strip()

        if not student_code or not name or not course or not email:
            messagebox.showerror(
                "Validation Error", "Please fill in all required fields."
            )
            return

        if not student_code.isalnum():
            messagebox.showerror(
                "Validation Error", "Student code must be alphanumeric."
            )
            return

        age = None
        if not age_str:
            if self.mode == "update" and self.student_data:
                age = self.student_data["age"]
            else:
                messagebox.showerror(
                    "Validation Error", "Please enter a valid age."
                )
                return
        else:
            try:
                age = int(age_str)
                if age < 5 or age > 100:
                    messagebox.showerror(
                        "Validation Error", "Age must be between 5 and 100."
                    )
                    return
            except ValueError:
                messagebox.showerror(
                    "Validation Error", "Please enter a valid age."
                )
                return

        student_dict = {
            "student_code": student_code,
            "name": name,
            "age": age,
            "course": course,
            "email": email,
        }

        temp_student = Student(
            student_code=student_code, name=name, age=age,
            course=course, email=email
        )
        errors = temp_student.validate()
        if errors:
            messagebox.showerror("Validation Error", "; ".join(errors))
            return

        if self.mode == "add":
            existing = find_student_by_id(student_code)
            if existing:
                messagebox.showerror(
                    "Duplicate", "Student code already exists."
                )
                return
            if not add_student(student_dict):
                messagebox.showerror(
                    "Error", "Failed to add student. Duplicate code?"
                )
                return
        elif self.mode == "update" and self.student_data:
            student_dict["age"] = age
            if not update_student(self.student_data["student_code"], student_dict):
                messagebox.showerror("Error", "Failed to update student.")
                return

        self.result = student_dict
        self.dialog.destroy()

    def _on_cancel(self):
        self.result = None
        self.dialog.destroy()


class StudentDetailsDialog:
    """Read-only dialog to view student details."""

    def __init__(self, parent, student_data):
        self.dialog = tk.Toplevel(parent)
        self.dialog.title("Student Details")
        self.dialog.resizable(False, False)
        self.dialog.transient(parent)
        self.dialog.grab_set()

        self._build_dialog(student_data)

        self.dialog.protocol("WM_DELETE_WINDOW", self._on_close)
        self.dialog.bind("<Escape>", lambda e: self._on_close())
        self.dialog.wait_window()

    def _build_dialog(self, data):
        frame = ttk.Frame(self.dialog, padding=20)
        frame.pack(fill=tk.BOTH)

        title_label = ttk.Label(frame, text="Student Details", font=("Segoe UI", 14, "bold"))
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 15))

        labels = ["Student Code", "Name", "Age", "Course", "Email"]
        for i, label in enumerate(labels):
            ttk.Label(frame, text=label + ":", font=("Segoe UI", 10, "bold")).grid(
                row=i + 1, column=0, sticky=tk.W, pady=6, padx=(0, 10)
            )
            value = data[label.lower().replace(" ", "_")]
            ttk.Label(frame, text=value, font=("Segoe UI", 10)).grid(
                row=i + 1, column=1, sticky=tk.W, pady=6
            )

        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=len(labels) + 1, column=0, columnspan=2, pady=20)
        ttk.Button(btn_frame, text="Close", command=self._on_close).pack()

    def _on_close(self):
        self.dialog.destroy()


class StudentManagementApp:
    """Main Tkinter application for Student Management System."""

    def __init__(self, root):
        self.root = root
        self.root.title("Student Management System")
        self.root.geometry("1000x650")
        self.root.minsize(900, 600)
        self.root.resizable(True, True)
        self._center_window()

        self.search_term = ""
        self.search_by = "Name"
        self.sort_column = None
        self.sort_reverse = False

        self._setup_styles()
        self._build_layout()
        self.refresh_students()
        
        # Bind database restore event
        self.root.bind("<<DatabaseRestored>>", self._on_database_restored)

    def _center_window(self):
        self.root.update_idletasks()
        w = self.root.winfo_width()
        h = self.root.winfo_height()
        sw = self.root.winfo_screenwidth()
        sh = self.root.winfo_screenheight()
        self.root.geometry(f"+{(sw - w) // 2}+{(sh - h) // 2}")

    def _setup_styles(self):
        style = ttk.Style()
        style.configure("Title.TLabel", font=("Segoe UI", 18, "bold"))
        style.configure("Subtitle.TLabel", font=("Segoe UI", 10), foreground="gray")
        style.configure("Card.TFrame", relief="solid", borderwidth=1)
        style.configure("CardValue.TLabel", font=("Segoe UI", 24, "bold"))
        style.configure("CardLabel.TLabel", font=("Segoe UI", 9), foreground="gray")
        style.configure("Action.TButton", font=("Segoe UI", 10), padding=8)
        style.configure("Treeview", font=("Segoe UI", 9), rowheight=28)
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("Status.TLabel", font=("Segoe UI", 9))

    def _build_layout(self):
        main_frame = ttk.Frame(self.root, padding=15)
        main_frame.pack(fill=tk.BOTH, expand=True)
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(4, weight=1)

        # HEADER
        header_frame = ttk.Frame(main_frame)
        header_frame.grid(row=0, column=0, sticky=tk.EW, pady=(0, 15))
        header_frame.columnconfigure(0, weight=1)

        ttk.Label(header_frame, text="Student Management System", style="Title.TLabel").grid(
            row=0, column=0, sticky=tk.W
        )
        ttk.Label(header_frame, text="Manage student records efficiently", style="Subtitle.TLabel").grid(
            row=1, column=0, sticky=tk.W
        )

        # DASHBOARD SUMMARY CARDS
        dashboard_frame = ttk.Frame(main_frame)
        dashboard_frame.grid(row=1, column=0, sticky=tk.EW, pady=(0, 15))
        dashboard_frame.columnconfigure((0, 1, 2), weight=1)

        self.card_total = self._create_card(dashboard_frame, "Total Students", "0", 0)
        self.card_courses = self._create_card(dashboard_frame, "Total Courses", "0", 1)
        self.card_avg_age = self._create_card(dashboard_frame, "Average Age", "0", 2)

        # ACTION AREA
        action_frame = ttk.Frame(main_frame)
        action_frame.grid(row=2, column=0, sticky=tk.EW, pady=(0, 10))

        buttons = [
            ("Add Student", self.open_add_dialog),
            ("Update", self.open_update_dialog),
            ("Delete", self.delete_student),
            ("View Details", self.view_details),
            ("Refresh", self.refresh_students),
            ("Export CSV", self.export_csv),
            ("Export Excel", self.export_excel),
            ("Backup / Restore", self.open_backup_restore),
        ]
        for i, (text, cmd) in enumerate(buttons):
            btn = ttk.Button(action_frame, text=text, command=cmd, style="Action.TButton")
            btn.pack(side=tk.LEFT, padx=4)

        # SEARCH AREA
        search_frame = ttk.Frame(main_frame)
        search_frame.grid(row=3, column=0, sticky=tk.EW, pady=(0, 10))
        search_frame.columnconfigure(1, weight=1)

        ttk.Label(search_frame, text="Search:", font=("Segoe UI", 10)).pack(side=tk.LEFT, padx=(0, 5))

        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var, width=35, font=("Segoe UI", 10))
        self.search_entry.pack(side=tk.LEFT, padx=5)
        self.search_entry.bind("<KeyRelease>", lambda e: self.search_students())
        self.search_entry.bind("<Return>", lambda e: self.search_students())

        ttk.Button(search_frame, text="Search", command=self.search_students).pack(side=tk.LEFT, padx=3)
        ttk.Button(search_frame, text="Clear", command=self.clear_search).pack(side=tk.LEFT, padx=3)

        ttk.Label(search_frame, text="Search By:", font=("Segoe UI", 10)).pack(side=tk.LEFT, padx=(15, 5))
        self.search_by_var = tk.StringVar(value="Name")
        search_combo = ttk.Combobox(
            search_frame, textvariable=self.search_by_var,
            values=["Name", "Student Code", "Course", "Email", "All Fields"],
            state="readonly", width=12
        )
        search_combo.pack(side=tk.LEFT, padx=3)
        search_combo.bind("<<ComboboxSelected>>", lambda e: self.search_students())

        # STUDENT TABLE with Course Statistics sidebar
        table_container = ttk.Frame(main_frame)
        table_container.grid(row=4, column=0, sticky=tk.NSEW)
        table_container.columnconfigure(0, weight=1)
        table_container.rowconfigure(0, weight=1)

        tree_frame = ttk.Frame(table_container)
        tree_frame.grid(row=0, column=0, sticky=tk.NSEW)
        table_container.columnconfigure(0, weight=1)
        table_container.rowconfigure(0, weight=1)

        columns = ("student_code", "name", "age", "course", "email")
        self.tree = ttk.Treeview(tree_frame, columns=columns, show="headings", selectmode="browse")

        for col in columns:
            self.tree.heading(col, text=col.replace("_", " ").title(),
                             command=lambda c=col: self._sort_treeview(c))
            self.tree.column(col, anchor=tk.CENTER if col in ("student_code", "age") else tk.W)

        self.tree.column("student_code", width=120, minwidth=100)
        self.tree.column("name", width=180, minwidth=120)
        self.tree.column("age", width=70, minwidth=60)
        self.tree.column("course", width=180, minwidth=120)
        self.tree.column("email", width=250, minwidth=180)

        v_scrollbar = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        h_scrollbar = ttk.Scrollbar(tree_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scrollbar.set, xscrollcommand=h_scrollbar.set)

        self.tree.grid(row=0, column=0, sticky=tk.NSEW)
        v_scrollbar.grid(row=0, column=1, sticky=tk.NS)
        h_scrollbar.grid(row=1, column=0, sticky=tk.EW)
        tree_frame.columnconfigure(0, weight=1)
        tree_frame.rowconfigure(0, weight=1)

        self.tree.bind("<Double-1>", lambda e: self.open_update_dialog())
        self.tree.bind("<Button-1>", self._on_tree_click)

        # Course Statistics Sidebar
        stats_frame = ttk.LabelFrame(table_container, text="Course Statistics", padding=10)
        stats_frame.grid(row=0, column=1, sticky=tk.NS, padx=(15, 0))
        stats_frame.rowconfigure(0, weight=1)

        self.stats_tree = ttk.Treeview(stats_frame, columns=("course", "count"), show="headings", height=15)
        self.stats_tree.heading("course", text="Course")
        self.stats_tree.heading("count", text="Count")
        self.stats_tree.column("course", width=150, anchor=tk.W)
        self.stats_tree.column("count", width=60, anchor=tk.CENTER)
        self.stats_tree.pack(fill=tk.BOTH, expand=True)

        # STATUS BAR
        status_frame = ttk.Frame(main_frame, relief="sunken", padding=(10, 5))
        status_frame.grid(row=5, column=0, sticky=tk.EW, pady=(10, 0))
        self.status_var = tk.StringVar(value="Ready")
        ttk.Label(status_frame, textvariable=self.status_var, style="Status.TLabel").pack(side=tk.LEFT)

    def _create_card(self, parent, label, value, column):
        card = ttk.Frame(parent, style="Card.TFrame", padding=15)
        card.grid(row=0, column=column, sticky=tk.EW, padx=5)
        card.columnconfigure(0, weight=1)

        ttk.Label(card, text=value, style="CardValue.TLabel").grid(row=0, column=0, sticky=tk.EW)
        ttk.Label(card, text=label, style="CardLabel.TLabel").grid(row=1, column=0, sticky=tk.EW, pady=(5, 0))

        if column == 0:
            self.card_total_value = card.children['!label']
        elif column == 1:
            self.card_courses_value = card.children['!label']
        else:
            self.card_avg_age_value = card.children['!label']

    def load_students(self):
        try:
            return get_all_students()
        except Exception:
            messagebox.showerror("Database Error", "Failed to load students from database.")
            return []

    def _update_dashboard(self, students):
        total = len(students)
        courses = len(set(s["course"] for s in students)) if students else 0
        avg_age = sum(int(s["age"]) for s in students) / len(students) if students else 0

        # Update card values
        self.card_total_value.config(text=str(total))
        self.card_courses_value.config(text=str(courses))
        self.card_avg_age_value.config(text=f"{avg_age:.1f}")

    def _update_course_stats(self):
        stats = get_course_statistics()
        for item in self.stats_tree.get_children():
            self.stats_tree.delete(item)
        for stat in stats:
            self.stats_tree.insert("", tk.END, values=(stat["course"], stat["count"]))

    def populate_tree(self, students):
        for item in self.tree.get_children():
            self.tree.delete(item)
        for s in students:
            self.tree.insert("", tk.END, values=(
                s["student_code"], s["name"], s["age"], s["course"], s["email"]
            ))

    def refresh_students(self):
        students = self.load_students()
        self.populate_tree(students)
        self._update_dashboard(students)
        self._update_course_stats()
        self.search_term = ""
        self.search_var.set("")
        self.search_by_var.set("Name")
        self.status_var.set(f"Total Students: {len(students)}")

    def search_students(self):
        term = self.search_var.get().strip()
        self.search_term = term
        self.search_by = self.search_by_var.get()

        if not term:
            students = self.load_students()
        else:
            search_by = self.search_by
            if search_by == "Name":
                students = search_students_by_name(term)
            elif search_by == "Student Code":
                students = search_students_by_code(term)
            elif search_by == "Course":
                students = search_students_by_course(term)
            elif search_by == "Email":
                students = search_students_by_email(term)
            elif search_by == "All Fields":
                students = search_students_all_fields(term)
            else:
                students = search_students_by_name(term)

        self.populate_tree(students)
        self._update_dashboard(students)
        total_all = len(self.load_students())
        if term:
            self.status_var.set(f"Showing {len(students)} of {total_all} students (filtered by {search_by}: '{term}')")
        else:
            self.status_var.set(f"Total Students: {len(students)}")

    def clear_search(self):
        self.search_var.set("")
        self.search_term = ""
        self.search_by_var.set("Name")
        self.refresh_students()

    def _get_visible_students(self):
        """Get the currently visible students from the Treeview."""
        students = []
        for item in self.tree.get_children():
            values = self.tree.item(item, "values")
            if values:
                students.append({
                    "student_code": values[0],
                    "name": values[1],
                    "age": values[2],
                    "course": values[3],
                    "email": values[4],
                })
        return students

    def _sort_treeview(self, column):
        if self.sort_column == column:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_column = column
            self.sort_reverse = False

        students = self._get_visible_students()
        col_index = {"student_code": 0, "name": 1, "age": 2, "course": 3, "email": 4}[column]

        def sort_key(s):
            val = s[list(s.keys())[col_index]]
            if column == "age":
                try:
                    return int(val)
                except (ValueError, TypeError):
                    return 0
            return str(val).lower()

        students.sort(key=sort_key, reverse=self.sort_reverse)
        self.populate_tree(students)

    def _on_tree_click(self, event):
        region = self.tree.identify("region", event.x, event.y)
        if region == "heading":
            return
        item = self.tree.identify_row(event.y)
        if item:
            self.tree.selection_set(item)

    def get_selected_student(self):
        selected = self.tree.selection()
        if not selected:
            return None
        values = self.tree.item(selected[0], "values")
        if not values:
            return None
        return {
            "student_code": values[0],
            "name": values[1],
            "age": values[2],
            "course": values[3],
            "email": values[4],
        }

    def open_add_dialog(self):
        dialog = StudentFormDialog(self.root, mode="add")
        if dialog.result:
            self.refresh_students()
            self.status_var.set("Student added successfully.")
            self.root.after(3000, lambda: self.status_var.set(f"Total Students: {len(self._get_visible_students())}"))

    def open_update_dialog(self):
        student = self.get_selected_student()
        if not student:
            messagebox.showinfo("No Selection", "Please select a student first.")
            return
        dialog = StudentFormDialog(self.root, mode="update", student_data=student)
        if dialog.result:
            self.refresh_students()
            self.status_var.set("Student updated successfully.")
            self.root.after(3000, lambda: self.status_var.set(f"Total Students: {len(self._get_visible_students())}"))

    def view_details(self):
        student = self.get_selected_student()
        if not student:
            messagebox.showinfo("No Selection", "Please select a student first.")
            return
        StudentDetailsDialog(self.root, student)

    def delete_student(self):
        student = self.get_selected_student()
        if not student:
            messagebox.showinfo("No Selection", "Please select a student first.")
            return
        code = student["student_code"]
        confirmed = messagebox.askyesno(
            "Confirm Delete", f"Are you sure you want to delete student {code}?"
        )
        if confirmed:
            if delete_student(code):
                self.refresh_students()
                self.status_var.set(f"Student {code} deleted successfully.")
                self.root.after(3000, lambda: self.status_var.set(f"Total Students: {len(self._get_visible_students())}"))
            else:
                messagebox.showerror("Error", "Failed to delete student.")

    def open_backup_restore(self):
        BackupRestoreDialog(self.root)

    def _on_database_restored(self, event):
        """Handle database restore event."""
        self.refresh_students()
        self.status_var.set("Database restored successfully.")

    def export_csv(self):
        """Export visible students to CSV file."""
        students = self._get_visible_students()
        if not students:
            messagebox.showinfo("No Data", "No students to export.")
            return

        filename = filedialog.asksaveasfilename(
            title="Export Students to CSV",
            defaultextension=".csv",
            initialfile="students.csv",
            filetypes=[("CSV Files", "*.csv"), ("All Files", "*.*")]
        )
        if not filename:
            return

        try:
            with open(filename, "w", newline="", encoding="utf-8-sig") as file:
                writer = csv.writer(file)
                writer.writerow(["Student Code", "Name", "Age", "Course", "Email"])
                for s in students:
                    writer.writerow([s["student_code"], s["name"], s["age"], s["course"], s["email"]])
            messagebox.showinfo("Success", "Students exported successfully to CSV.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export CSV: {e}")

    def export_excel(self):
        """Export visible students to Excel file."""
        if not HAS_OPENPYXL:
            messagebox.showerror(
                "Missing Dependency",
                "openpyxl is not installed.\n\n"
                "Install it using:\n"
                "python -m pip install openpyxl"
            )
            return

        students = self._get_visible_students()
        if not students:
            messagebox.showinfo("No Data", "No students to export.")
            return

        filename = filedialog.asksaveasfilename(
            title="Export Students to Excel",
            defaultextension=".xlsx",
            initialfile="students.xlsx",
            filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
        )
        if not filename:
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter

            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Students"

            headers = ["Student Code", "Name", "Age", "Course", "Email"]
            for col_idx, header in enumerate(headers, 1):
                cell = worksheet.cell(row=1, column=col_idx, value=header)
                cell.font = Font(bold=True)

            for row_idx, s in enumerate(students, 2):
                worksheet.cell(row=row_idx, column=1, value=s["student_code"])
                worksheet.cell(row=row_idx, column=2, value=s["name"])
                worksheet.cell(row=row_idx, column=3, value=s["age"])
                worksheet.cell(row=row_idx, column=4, value=s["course"])
                worksheet.cell(row=row_idx, column=5, value=s["email"])

            for col_idx in range(1, 6):
                column_letter = get_column_letter(col_idx)
                max_length = len(headers[col_idx - 1])
                for row_idx in range(2, len(students) + 2):
                    cell_value = str(worksheet.cell(row=row_idx, column=col_idx).value or "")
                    max_length = max(max_length, len(cell_value))
                worksheet.column_dimensions[column_letter].width = max_length + 2

            worksheet.freeze_panes = "A2"

            worksheet.auto_filter.ref = f"A1:E{len(students) + 1}"

            workbook.save(filename)
            messagebox.showinfo("Success", "Students exported successfully to Excel.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export Excel: {e}")


def show_login(root):
    """Show login dialog and return (success, username)."""
    try:
        login = LoginDialog(root)
        return login.result, login.username
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Failed to create login dialog: {e}")
        return False, ""


def main():
    try:
        initialize_database()
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Failed to initialize database: {e}")
        return

    root = tk.Tk()
    root.withdraw()

    try:
        if not has_admin_accounts():
            setup = AdminSetupDialog(root)
            if not setup.result:
                root.destroy()
                return

        success, username = show_login(root)
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Login dialog error: {e}")
        root.destroy()
        return

    if not success:
        root.destroy()
        return

    try:
        root.deiconify()
        app = StudentManagementApp(root)
        root.mainloop()
    except Exception as e:
        import traceback
        traceback.print_exc()
        print(f"Application error: {e}")
        root.destroy()


if __name__ == "__main__":
    main()
